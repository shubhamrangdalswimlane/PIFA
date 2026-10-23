"""files-svc Excel sandbox and gateway (local storage backend, auth disabled)."""

import importlib
import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "files-svc"))


def xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE", "local")
    monkeypatch.setenv("LOCAL_STORAGE_DIR", str(tmp_path))
    monkeypatch.setenv("AUTH_DISABLED", "true")
    monkeypatch.setenv("POLL_SECONDS", "0")
    import app.main

    importlib.reload(app.main)
    return TestClient(app.main.app), tmp_path, app.main


def test_valid_forecast_is_processed(client):
    c, root, _ = client
    r = c.post("/ingest/forecast", data={"uploadedBy": "sm.north"},
               files={"file": ("f.xlsx", xlsx([["Product Name", "Forecast Quantity"], ["Printer Model A", 25]]))}).json()
    assert r["status"] == "processed" and r["rows"] == [{"row": 2, "product": "Printer Model A", "values": {"qty": 25}}]
    assert list((root / "files/processed").iterdir())


def test_bad_rows_reject_file_with_reason(client):
    c, root, _ = client
    content = xlsx([["Product Name", "Quantity"], ["A", -1], ["B", 2.5], ["A", 3]])
    r = c.post("/ingest/forecast", files={"file": ("f.xlsx", content)}).json()
    assert r["status"] == "rejected" and len(r["errors"]) == 3 and r["rows"] == []
    assert any(p.name.endswith(".reason.txt") for p in (root / "files/rejected").iterdir())


def test_allocation_needs_region_columns(client):
    c, _, _ = client
    r = c.post("/ingest/allocation", files={"file": ("a.xlsx", xlsx([["Product Name", "North"], ["A", 1]]))}).json()
    assert r["status"] == "rejected" and "East" in r["errors"][0]["message"]


def test_non_xlsx_and_macro_files_rejected(client):
    c, _, _ = client
    assert c.post("/ingest/forecast", files={"file": ("f.csv", b"Product Name,Qty\nA,1")}).json()["status"] == "rejected"
    assert c.post("/ingest/forecast", files={"file": ("f.xlsx", b"not a zip")}).json()["status"] == "rejected"


def test_export_neutralises_formulas(client):
    c, root, _ = client
    r = c.post("/exports", json={"filename": "x.xlsx", "sheet": "S", "columns": ["Product Name"], "rows": [["=HYPERLINK(1)"]]})
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A2"].value == "'=HYPERLINK(1)"
    assert list((root / "files/outbound").iterdir())


def test_poller_moves_upstream_files(client):
    c, root, main = client
    assert c.post("/upstream/inventory", files={"file": ("inv.xlsx", xlsx([["Product Name", "Available Inventory"], ["A", 5]]))}).status_code == 200
    assert main.poll_once() == 1
    assert not list((root / "files/inbound/inventory").iterdir())
    assert list((root / "files/processed").iterdir())


def test_product_name_header_variants_and_month_column(client):
    c, _, _ = client
    content = xlsx([["Product_Name", "Month", "Forecast_Quantity"],
                    ["Printer Model A", "Nov-2026", 10], ["Printer Model A", "2026-12", 12], ["New Model Z", "12/2026", 3]])
    r = c.post("/ingest/forecast", files={"file": ("m.xlsx", content)}).json()
    assert r["status"] == "processed", r["errors"]
    assert [(x["product"], x["month"], x["values"]["qty"]) for x in r["rows"]] == [
        ("Printer Model A", "2026-11", 10), ("Printer Model A", "2026-12", 12), ("New Model Z", "2026-12", 3)]


def test_month_column_errors_and_duplicates_per_month(client):
    c, _, _ = client
    content = xlsx([["Product Name", "Month", "Quantity"], ["A", "Foo-26", 1], ["B", "2026-11", 1], ["B", "Nov 2026", 2]])
    r = c.post("/ingest/forecast", files={"file": ("m.xlsx", content)}).json()
    assert r["status"] == "rejected" and len(r["errors"]) == 2
