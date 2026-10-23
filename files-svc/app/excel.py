"""Excel sandbox: validate and parse uploaded .xlsx files into rows (values only, no formulas or macros)."""

import io
import zipfile
import re
from datetime import UTC, date, datetime

from openpyxl import Workbook, load_workbook

REGIONS = ["North", "East", "West", "South"]
MAX_BYTES = 4 * 1024 * 1024
MAX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ROWS = 2000

# kind -> list of (value key, accepted header names)
COLUMNS = {
    "forecast": [("qty", ["forecast quantity", "quantity", "forecast qty", "qty"])],
    "inventory": [("qty", ["available inventory", "inventory", "quantity", "qty"])],
    "allocation": [(r, [r.lower()]) for r in REGIONS],
}


class FileRejected(Exception):
    def __init__(self, message: str, errors: list[dict] | None = None):
        """Signal that an uploaded file is rejected as a whole.

        Args:
            message: Human-readable reason for the rejection.
            errors: Row errors as `{"row": int, "message": str}` dicts; defaults to a single row-0 error
                carrying `message`.

        Returns:
            None
        """
        super().__init__(message)
        self.errors = errors or [{"row": 0, "message": message}]


def check_container(filename: str, content: bytes) -> None:
    """Size, extension, zip magic bytes, zip-bomb and macro checks before openpyxl touches the file.

    Args:
        filename: Uploaded file name; must end in ".xlsx".
        content: The raw file bytes.

    Returns:
        None when the file passes every check.

    Raises:
        FileRejected: When the file is not .xlsx, is over 4 MB, is not a zip, unzips to more than 50 MB,
            contains macros (vbaProject), is not an Excel workbook, or is a corrupt zip.
    """
    if not filename.lower().endswith(".xlsx"):
        raise FileRejected("Only .xlsx files are accepted")
    if len(content) > MAX_BYTES:
        raise FileRejected("File larger than 4 MB")
    if not content.startswith(b"PK\x03\x04"):
        raise FileRejected("Not a valid .xlsx (zip) file")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            names = z.namelist()
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise FileRejected("File expands too much when unzipped")
            if any("vbaproject" in n.lower() for n in names):
                raise FileRejected("Files with macros are not accepted")
            if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
                raise FileRejected("Not an Excel workbook")
    except zipfile.BadZipFile as e:
        raise FileRejected("Corrupt .xlsx file") from e


def _whole(value) -> int | None:
    """Read a cell value as a whole number (quantity), accepting ints, whole floats and digit strings.

    Args:
        value: The raw cell value.

    Returns:
        The integer value, or None when it is empty, a boolean, fractional or not a number.
    """
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().lstrip("-").isdigit():
        return int(value.strip())
    return None


PRODUCT_HEADERS = ["product name", "productname", "product"]
MONTH_HEADERS = ["month", "month year", "monthyear", "period", "forecast month"]
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _norm_header(h) -> str:
    """'Product_Name', 'product-name ', 'PRODUCT NAME' -> 'product name'.

    Args:
        h: A header cell value.

    Returns:
        The header lower-cased, with "_" and "-" turned into spaces and whitespace collapsed.
    """
    return " ".join(re.sub(r"[_\-]+", " ", str(h)).lower().split())


def parse_month(value) -> str:
    """Month-year cell -> 'YYYY-MM'. Accepts dates and 2026-11, 11/2026, Nov-2026, Nov 2026, November 2026, Nov-26.

    Args:
        value: The Month cell value (date/datetime or text).

    Returns:
        The month as "YYYY-MM", e.g. "2026-11".

    Raises:
        ValueError: When the cell is empty, not a recognised month-year, or outside months 1-12 / years 2000-2100.
    """
    if isinstance(value, datetime | date):
        return f"{value.year:04d}-{value.month:02d}"
    s = str(value or "").strip().lower()
    if not s:
        raise ValueError("empty (use e.g. 2026-11 or Nov-2026)")
    if m := re.fullmatch(r"(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?(?:[ t].*)?", s):
        year, month = int(m.group(1)), int(m.group(2))
    elif m := re.fullmatch(r"(\d{1,2})[-/.](\d{4})", s):
        year, month = int(m.group(2)), int(m.group(1))
    elif (m := re.fullmatch(r"([a-z]{3})[a-z]*[\s\-/.']*(\d{2}|\d{4})", s)) and m.group(1) in _MONTHS:
        month, year = _MONTHS[m.group(1)], int(m.group(2))
        year += 2000 if year < 100 else 0
    else:
        raise ValueError(f"'{value}' is not a month-year (use e.g. 2026-11 or Nov-2026)")
    if not (1 <= month <= 12 and 2000 <= year <= 2100):
        raise ValueError(f"'{value}' is not a valid month-year")
    return f"{year:04d}-{month:02d}"


def parse(kind: str, content: bytes) -> tuple[list[dict], list[dict]]:
    """Return (rows, errors). Rows: {"row", "product", "values"} plus "month" ("YYYY-MM") when the file has a
    Month column, so one file can carry several month-years. Any error rejects the whole file.

    Only the first sheet is read, as values (no formulas). The product column may be "Product Name" or
    "Product_Name"; empty quantity cells count as 0; duplicate products (per month) are errors.

    Args:
        kind: File kind: "forecast", "inventory" or "allocation" (decides the quantity columns).
        content: The raw .xlsx bytes (already passed `check_container`).

    Returns:
        A `(rows, errors)` tuple: rows are `{"row": int, "product": str, "values": {key: int}[, "month": "YYYY-MM"]}`
        with `values` keyed `qty` (forecast/inventory) or by region (allocation); errors are `{"row": int, "message": str}`.

    Raises:
        FileRejected: For an unknown kind, an unreadable workbook, an empty first sheet, a missing product or
            quantity column, or more than MAX_ROWS data rows.
    """
    if kind not in COLUMNS:
        raise FileRejected(f"Unknown file kind {kind}")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - any parser failure means a bad file
        raise FileRejected(f"Cannot read workbook: {e}") from e
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise FileRejected("The first sheet is empty")
    index = {_norm_header(h): i for i, h in enumerate(header) if h is not None}
    name_col = next((index[n] for n in PRODUCT_HEADERS if n in index), None)
    if name_col is None:
        raise FileRejected("Missing 'Product Name' (or 'Product_Name') column")
    month_col = next((index[n] for n in MONTH_HEADERS if n in index), None)
    cols = {}
    for key, names in COLUMNS[kind]:
        found = next((index[n] for n in names if n in index), None)
        if found is None:
            raise FileRejected(f"Missing '{names[0].title()}' column")
        cols[key] = found

    rows, errors, seen = [], [], {}
    for n, raw in enumerate(rows_iter, start=2):
        if n - 1 > MAX_ROWS:
            raise FileRejected(f"More than {MAX_ROWS} rows")
        name = raw[name_col] if name_col < len(raw) else None
        if name is None or str(name).strip() == "":
            if any(v not in (None, "") for v in raw):
                errors.append({"row": n, "message": "Product Name is empty"})
            continue
        product = " ".join(str(name).split())
        month = None
        if month_col is not None:
            cell = raw[month_col] if month_col < len(raw) else None
            try:
                month = parse_month(cell)
            except ValueError as e:
                errors.append({"row": n, "message": f"Month: {e}"})
                continue
        dup = (product.lower(), month)
        if dup in seen:
            where = f" for {month}" if month else ""
            errors.append({"row": n, "message": f"Duplicate product '{product}'{where} (first on row {seen[dup]})"})
            continue
        seen[dup] = n
        values = {}
        for key, i in cols.items():
            cell = raw[i] if i < len(raw) else None
            qty = _whole(0 if cell in (None, "") else cell)
            if qty is None or qty < 0:
                errors.append({"row": n, "message": f"{key if key != 'qty' else 'Quantity'} must be a whole number of 0 or more (got {cell!r})"})
            else:
                values[key] = qty
        if len(values) == len(cols):
            row = {"row": n, "product": product, "values": values}
            if month_col is not None:
                row["month"] = month
            rows.append(row)
    wb.close()
    return rows, errors


def _safe(value):
    """Neutralise spreadsheet formula injection in exported text cells.

    Args:
        value: A header or cell value to export.

    Returns:
        The value prefixed with "'" when it is text starting with =, +, -, @, tab or carriage return;
        otherwise the value unchanged.
    """
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def build(sheet: str, columns: list[str], rows: list[list]) -> bytes:
    """Build an export workbook: one sheet with a header row and data rows, formula-injection safe.

    The header row is frozen and column A widened for product names.

    Args:
        sheet: Sheet name (cut to Excel's 31-character limit; "Sheet1" if empty).
        columns: Header row labels.
        rows: Data rows, one list of cell values per row.

    Returns:
        The .xlsx file as bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31] or "Sheet1"
    ws.append([_safe(c) for c in columns])
    for row in rows:
        ws.append([_safe(v) for v in row])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 32
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stamp() -> str:
    """Current UTC time as a compact timestamp used to prefix stored file names.

    Returns:
        A string like "20261022T100000" (YYYYMMDDTHHMMSS).
    """
    return datetime.now(UTC).strftime("%Y%m%dT%H%M%S")
